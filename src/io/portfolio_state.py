from dataclasses import dataclass
from pathlib import Path
import json

import pandas as pd
from openpyxl import load_workbook

from src.domain.asset import PriceSnapshot
from src.domain.portfolio import Portfolio
from src.domain.trade import Trade


@dataclass
class PortfolioStateManager:
    """Persists and loads portfolio state (cash + positions) between sessions.

    State is stored as JSON with the current snapshot, a history of
    daily VL values, daily positions, and all executed operations.
    """

    state_dir: str = "outputs/state"

    @property
    def _state_path(self) -> Path:
        return Path(self.state_dir) / "portfolio_state.json"

    def save(
        self,
        portfolio: Portfolio,
        prices: PriceSnapshot,
        date: pd.Timestamp,
        trades: list[Trade] | None = None,
        strategy_name: str = "merton_custom",
    ) -> None:
        """Save current portfolio state to disk."""
        Path(self.state_dir).mkdir(parents=True, exist_ok=True)

        state = self._load_raw() if self._state_path.exists() else {
            "history": [],
            "positions_history": [],
            "operations_history": [],
        }

        total_value = portfolio.total_value(prices)

        snapshot = {
            "date": date.isoformat(),
            "cash": portfolio.cash,
            "positions": portfolio.positions,
            "total_value": total_value,
        }
        state["current"] = snapshot

        # VL history (dedup by date)
        history = state.get("history", [])
        history = [h for h in history if h["date"] != date.isoformat()]
        history.append({"date": date.isoformat(), "total_value": total_value, "cash": portfolio.cash})
        history.sort(key=lambda h: h["date"])
        state["history"] = history

        # Positions history (daily snapshot of all positions with values)
        pos_history = state.get("positions_history", [])
        pos_history = [p for p in pos_history if p["date"] != date.isoformat()]
        for ticker, shares in portfolio.positions.items():
            if shares != 0:
                price = prices.get_price(ticker)
                value = shares * price
                weight = value / total_value if total_value else 0.0
                pos_history.append({
                    "date": date.isoformat(),
                    "ticker": ticker,
                    "shares": shares,
                    "price": price,
                    "value": value,
                    "weight": weight,
                })
        pos_history.sort(key=lambda p: (p["date"], p["ticker"]))
        state["positions_history"] = pos_history

        # Operations history (all trades)
        ops_history = state.get("operations_history", [])
        if trades:
            date_str = date.isoformat()
            # Remove any operations for today (idempotent re-runs)
            ops_history = [o for o in ops_history if o["date"] != date_str]
            for t in trades:
                is_liquidation = getattr(t.order.signal, "reason", "") == "liquidation"
                ops_history.append({
                    "date": date_str,
                    "ticker": t.ticker,
                    "shares": t.shares,
                    "price": t.price,
                    "value": t.shares * t.price,
                    "cost": t.cost,
                    "direction": "buy" if t.shares > 0 else "sell",
                    "strategy": "transicion" if is_liquidation else strategy_name,
                })
            ops_history.sort(key=lambda o: (o["date"], o["ticker"]))
        state["operations_history"] = ops_history

        self._state_path.write_text(json.dumps(state, indent=2, default=str))

    def load(self) -> dict | None:
        """Load last saved state. Returns None if no state exists."""
        if not self._state_path.exists():
            return None
        state = self._load_raw()
        return state.get("current")

    def load_into_portfolio(self, portfolio: Portfolio) -> pd.Timestamp | None:
        """Restore a portfolio from saved state. Returns the date of saved state."""
        current = self.load()
        if current is None:
            return None
        portfolio._cash = current["cash"]
        portfolio._positions = {k: v for k, v in current["positions"].items()}
        return pd.Timestamp(current["date"])

    def value_history(self) -> pd.Series:
        """Return the VL time series from saved history."""
        if not self._state_path.exists():
            return pd.Series(dtype=float)
        state = self._load_raw()
        history = state.get("history", [])
        if not history:
            return pd.Series(dtype=float)
        dates = [pd.Timestamp(h["date"]) for h in history]
        values = [h["total_value"] for h in history]
        return pd.Series(values, index=dates, name="total_value")

    def positions_history(self) -> pd.DataFrame:
        """Return all daily position snapshots as a DataFrame."""
        if not self._state_path.exists():
            return pd.DataFrame()
        state = self._load_raw()
        records = state.get("positions_history", [])
        if not records:
            return pd.DataFrame()
        df = pd.DataFrame(records)
        df["date"] = pd.to_datetime(df["date"])
        return df

    def operations_history(self) -> pd.DataFrame:
        """Return all historical trades as a DataFrame."""
        if not self._state_path.exists():
            return pd.DataFrame()
        state = self._load_raw()
        records = state.get("operations_history", [])
        if not records:
            return pd.DataFrame()
        df = pd.DataFrame(records)
        df["date"] = pd.to_datetime(df["date"])
        return df

    def _load_raw(self) -> dict:
        return json.loads(self._state_path.read_text())

    # ── estado.json ───────────────────────────────────────

    @property
    def _estado_path(self) -> Path:
        return Path(self.state_dir) / "estado.json"

    def load_estado(self) -> dict | None:
        """Read estado.json. Returns the dict or None if missing/uninitialised."""
        if not self._estado_path.exists():
            return None
        data = json.loads(self._estado_path.read_text())
        if not data.get("iniciado", False):
            return None
        return data

    def save_estado(
        self,
        portfolio: Portfolio,
        prices: PriceSnapshot,
        date: pd.Timestamp,
        costes_acumulados: float,
        n_operaciones: int,
        strategy_name: str = "merton_custom",
        capital_inicial: float = 10_000_000.0,
        fecha_inicio: str | None = None,
    ) -> None:
        """Write estado.json with the current portfolio state."""
        Path(self.state_dir).mkdir(parents=True, exist_ok=True)

        valor_cartera = portfolio.total_value(prices)

        # Preserve peak from previous estado
        prev = self.load_estado()
        prev_peak = prev.get("peak_valor", 0.0) if prev else 0.0
        peak = max(prev_peak, valor_cartera)

        if fecha_inicio is None:
            if prev and prev.get("fecha_inicio"):
                fecha_inicio = prev["fecha_inicio"]
            else:
                fecha_inicio = date.strftime("%Y-%m-%d")

        posiciones = {
            t: s for t, s in portfolio.positions.items() if abs(s) > 1e-9
        }

        state = {
            "iniciado": True,
            "estrategia": strategy_name,
            "fecha_inicio": fecha_inicio,
            "capital_inicial": capital_inicial,
            "valor_cartera": valor_cartera,
            "posiciones": posiciones,
            "cash": portfolio.cash,
            "peak_valor": peak,
            "n_operaciones": n_operaciones,
            "costes_acumulados": costes_acumulados,
            "fecha_ultima_ejecucion": date.strftime("%Y-%m-%d"),
        }
        self._estado_path.write_text(json.dumps(state, indent=4, default=str))

    # ── Historial_Grupo4.xlsx ──────────────────────────────

    _HISTORIAL_HEADERS = [
        "Fecha",
        "Valor Cartera",
        "Retorno diario",
        "Retorno acum.",
        "N posiciones",
        "Importe operado",
        "Coste (EUR)",
        "N ops dia",
        "N ops acum.",
        "Costes acum.",
        "Cash",
        "Decision",
    ]

    def _historial_title(self, strategy_name: str) -> str:
        return f"HISTORIAL SIMULACION — GRUPO 4 | Estrategia: {strategy_name} | Multi-asset"

    def save_historial(
        self,
        date: pd.Timestamp,
        valor_cartera: float,
        retorno_diario: float,
        retorno_acum: float,
        n_posiciones: int,
        importe_operado: float,
        coste_dia: float,
        n_ops_dia: int,
        n_ops_acum: int,
        costes_acum: float,
        cash: float,
        decision: str,
        group_name: str = "Grupo4",
        strategy_name: str = "merton_custom",
    ) -> Path:
        """Append a row to Historial_Grupo4.xlsx, creating new section if needed."""
        historial_path = Path(self.state_dir) / f"Historial_{group_name}.xlsx"
        title = self._historial_title(strategy_name)

        if historial_path.exists():
            wb = load_workbook(historial_path)
            ws = wb.active
        else:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "Historial"

        # Find if the new section already exists
        new_section_row = None
        date_str = date.strftime("%Y-%m-%d")
        for row_idx in range(1, ws.max_row + 1):
            cell_val = ws.cell(row=row_idx, column=1).value
            if cell_val and str(cell_val).startswith(title):
                new_section_row = row_idx
                break

        if new_section_row is None:
            # Add separator + new section
            next_row = ws.max_row + 2  # blank separator row
            ws.cell(row=next_row, column=1, value=title)
            next_row += 1
            for col_idx, header in enumerate(self._HISTORIAL_HEADERS, 1):
                ws.cell(row=next_row, column=col_idx, value=header)
            data_start_row = next_row + 1
        else:
            # Section exists — find header row (next row after title)
            header_row = new_section_row + 1
            data_start_row = header_row + 1

            # Dedup: remove any existing row for this date
            row_idx = data_start_row
            while row_idx <= ws.max_row:
                cell_val = ws.cell(row=row_idx, column=1).value
                if cell_val is not None:
                    cell_str = str(cell_val)[:10]
                    if cell_str == date_str:
                        ws.delete_rows(row_idx)
                        continue
                row_idx += 1
            data_start_row = ws.max_row + 1

        # Append the new row
        row_data = [
            date.to_pydatetime(),
            round(valor_cartera, 2),
            round(retorno_diario, 6),
            round(retorno_acum, 6),
            n_posiciones,
            round(importe_operado, 2),
            round(coste_dia, 2),
            n_ops_dia,
            n_ops_acum,
            round(costes_acum, 2),
            round(cash, 2),
            decision,
        ]
        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=data_start_row, column=col_idx, value=value)

        wb.save(historial_path)
        return historial_path

    # ── Legacy data extraction (E4 → new strategy migration) ──────

    def load_legacy_from_historial(
        self, group_name: str = "Grupo4",
    ) -> tuple[list[dict], list[dict]]:
        """Extract VL history and operations from the old Historial section.

        Returns:
            (vl_entries, legacy_operations) where each vl_entry is
            {date, total_value} and each operation is a dict compatible
            with operations_history format + strategy='E4'.
        """
        path = Path(self.state_dir) / f"Historial_{group_name}.xlsx"
        if not path.exists():
            return [], []

        df = pd.read_excel(path, header=1)
        df = df[df["Fecha"].notna()].copy()
        if df.empty:
            return [], []
        df["Fecha"] = pd.to_datetime(df["Fecha"])

        # VL entries — one per day
        vl_entries = [
            {"date": row["Fecha"], "total_value": float(row["Valor cartera"])}
            for _, row in df.iterrows()
        ]

        # Operations — only days with actual trades (COMPRAR / VENDER)
        iuse_ct = 0.00021
        operations = []
        for _, row in df.iterrows():
            decision = str(row["Decision"]).strip().upper()
            if decision == "MANTENER":
                continue

            importe = float(row.get("Importe (EUR)", 0) or 0)
            coste = float(row.get("Coste (EUR)", 0) or 0)
            precio = float(row["Precio ETF"])
            date_str = row["Fecha"].isoformat()

            if importe <= 0:
                continue

            if decision == "COMPRAR":
                shares_approx = importe / (precio * (1 + iuse_ct))
                operations.append({
                    "date": date_str,
                    "ticker": "IUSE.L",
                    "shares": shares_approx,
                    "price": precio,
                    "value": importe,
                    "cost": coste,
                    "direction": "buy",
                    "strategy": "E4",
                })
            elif decision == "VENDER":
                shares_approx = importe / (precio * (1 - iuse_ct))
                operations.append({
                    "date": date_str,
                    "ticker": "IUSE.L",
                    "shares": -shares_approx,
                    "price": precio,
                    "value": importe,
                    "cost": coste,
                    "direction": "sell",
                    "strategy": "E4",
                })

        return vl_entries, operations

    def seed_legacy_data(
        self,
        vl_entries: list[dict],
        legacy_ops: list[dict],
        vl_tracker,
    ) -> None:
        """Write legacy VL + operations into state files.

        Called once during migration from the old strategy.
        """
        # Seed VL history
        for entry in vl_entries:
            vl_tracker.record(entry["date"], entry["total_value"])

        # Seed legacy operations into portfolio_state.json
        if not legacy_ops or not self._state_path.exists():
            return

        state = self._load_raw()
        existing_ops = state.get("operations_history", [])

        # Tag existing ops that don't have a strategy field
        for op in existing_ops:
            if "strategy" not in op:
                op["strategy"] = "unknown"

        # Merge: legacy first, then new (avoid duplicates by date+ticker)
        existing_keys = {(o["date"], o["ticker"]) for o in existing_ops}
        for op in legacy_ops:
            key = (op["date"], op["ticker"])
            if key not in existing_keys:
                existing_ops.append(op)

        existing_ops.sort(key=lambda o: (o["date"], o.get("strategy", ""), o["ticker"]))
        state["operations_history"] = existing_ops

        # Also seed VL into history array
        existing_vl = state.get("history", [])
        existing_dates = {h["date"] for h in existing_vl}
        for entry in vl_entries:
            d = entry["date"].isoformat() if hasattr(entry["date"], "isoformat") else str(entry["date"])
            if d not in existing_dates:
                existing_vl.append({
                    "date": d,
                    "total_value": entry["total_value"],
                    "cash": 0.0,
                })
        existing_vl.sort(key=lambda h: h["date"])
        state["history"] = existing_vl

        self._state_path.write_text(json.dumps(state, indent=2, default=str))
