"""
simulacion_diaria.py
====================
Script de seguimiento diario para la simulación en tiempo real.
El Programador de Tareas de Windows lo ejecuta automáticamente cada día hábil a las 21:00h.

Configuración inicial (cambiar antes de empezar):
    - ETF_TICKER, ESTRATEGIA, GRUPO, CAPITAL_INICIAL
    - EMAIL_REMITENTE, EMAIL_PASSWORD (ver instrucciones abajo)
"""

import os
import sys
import json
import warnings
import numpy as np
import pandas as pd
import yfinance as yf
from datetime import datetime, date
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders

warnings.filterwarnings('ignore')

# ===========================================================================
# CONFIGURACION — CAMBIAR ANTES DE EMPEZAR
# ===========================================================================

ETF_TICKER       = "IUSE.L"              # Ticker del ETF elegido
ESTRATEGIA       = "E4"                  # E1, E2, E3, E4 o E5
GRUPO            = "X"                   # Numero de vuestro grupo
CAPITAL_INICIAL  = 10_000_000.0          # Capital inicial en EUR

# Email
EMAIL_REMITENTE  = "vuestro@outlook.com"    # CAMBIAR
EMAIL_PASSWORD   = "vuestra_contrasena"     # CAMBIAR (ver instrucciones abajo)
EMAIL_DESTINO    = "aguilabert@afi.es"
EMAIL_CC         = "vuestro@outlook.com"    # Copia a vosotros mismos (recomendado)

# INSTRUCCIONES PARA EMAIL_PASSWORD:
# Outlook NO permite usar tu contrasena normal por SMTP.
# Debes crear una "contrasena de aplicacion":
#   1. Ve a https://account.microsoft.com/security
#   2. Activa verificacion en dos pasos si no la tienes
#   3. Busca "Contrasenas de aplicacion" y crea una nueva
#   4. Copia aqui la contrasena generada (sin espacios)

# Costes de transaccion calculados con Bloomberg enero 2026
CT = {
    "MSE.PA":  0.001738,
    "IUSE.L":  0.000210,
    "IEMA.L":  0.000491,
    "IUSN.DE": 0.000605,
}

# Parametros del modelo (deben coincidir con config.py)
GAMMA         = -2
EWMA_LAMBDA   = 0.94
MU_SHRINKAGE  = 0.15
SIGMA_TARGET  = 0.15
MAX_LEVERAGE  = 1.0
WARMUP_DAYS   = 252

# Parametros E4/E5
SMA_WINDOW         = 200
DD_ENTRY_THRESHOLD = -0.08
DD_EXIT_THRESHOLD  = -0.03

# Rutas
BASE_DIR    = Path(__file__).parent
ESTADO_FILE = BASE_DIR / "simulacion_estado.json"
LOG_FILE    = BASE_DIR / "simulacion_log.csv"
OUTPUT_DIR  = BASE_DIR / "Simulacion_real"


# ===========================================================================
# ESTIMACION DE PARAMETROS
# ===========================================================================

def estimar_parametros(precios, rf_anual):
    retornos = np.log(precios / precios.shift(1)).dropna()
    if len(retornos) < WARMUP_DAYS:
        return None

    var_ewma = retornos.iloc[0] ** 2
    for r in retornos:
        var_ewma = EWMA_LAMBDA * var_ewma + (1 - EWMA_LAMBDA) * r ** 2
    sigma_anual = np.sqrt(var_ewma * 252)

    mu_historico = retornos.mean() * 252
    mu_estimado  = (1 - MU_SHRINKAGE) * mu_historico + MU_SHRINKAGE * 0.07

    rra = 1 - GAMMA
    if sigma_anual <= 0:
        return None

    alpha_star = np.clip((mu_estimado - rf_anual) / (rra * sigma_anual ** 2),
                         0.0, MAX_LEVERAGE)

    ct    = CT.get(ETF_TICKER, 0.001)
    ancho = np.sqrt(3 * ct * alpha_star / (rra * sigma_anual ** 2)) if alpha_star > 0 else 0.05

    return {
        "mu": mu_estimado, "sigma": sigma_anual,
        "alpha_star": alpha_star,
        "banda_inf": max(0.0, alpha_star - ancho),
        "banda_sup": min(MAX_LEVERAGE, alpha_star + ancho),
        "rf": rf_anual,
    }


# ===========================================================================
# LOGICA DE ESTRATEGIAS
# ===========================================================================

def decidir_alpha(params, precios, alpha_actual, peak_valor, valor_cartera, estado):
    alpha_star = params["alpha_star"]
    sigma      = params["sigma"]
    banda_inf  = params["banda_inf"]
    banda_sup  = params["banda_sup"]
    decision   = "MANTENER"
    alpha_nuevo = alpha_actual
    razon       = ""

    def evaluar_banda(objetivo):
        nonlocal decision, alpha_nuevo, razon, banda_inf, banda_sup
        ancho = banda_sup - alpha_star
        banda_inf = max(0.0, objetivo - ancho)
        banda_sup = min(MAX_LEVERAGE, objetivo + ancho)
        if alpha_actual < banda_inf:
            decision    = "COMPRAR"
            alpha_nuevo = objetivo
            razon       = f"a={alpha_actual:.3f} < banda_inf={banda_inf:.3f}"
        elif alpha_actual > banda_sup:
            decision    = "VENDER"
            alpha_nuevo = objetivo
            razon       = f"a={alpha_actual:.3f} > banda_sup={banda_sup:.3f}"

    if ESTRATEGIA == "E1":
        evaluar_banda(alpha_star)

    elif ESTRATEGIA == "E2":
        mom = 0.0
        if len(precios) >= 252:
            mom = (precios.iloc[-1] / precios.iloc[-252] - 1) - \
                  (precios.iloc[-1] / precios.iloc[-21]  - 1)
        alpha_adj = np.clip(alpha_star * (1 + 0.5 * np.sign(mom)), 0, MAX_LEVERAGE)
        evaluar_banda(alpha_adj)
        if razon:
            razon += f" | momentum={mom:.2%}"

    elif ESTRATEGIA == "E3":
        alpha_vol = min(SIGMA_TARGET / sigma, MAX_LEVERAGE) if sigma > 0 else alpha_star
        evaluar_banda(alpha_vol)
        if razon:
            razon += f" | sigma={sigma:.1%} alpha_vol={alpha_vol:.3f}"

    elif ESTRATEGIA == "E4":
        alpha_vol = min(SIGMA_TARGET / sigma, MAX_LEVERAGE) if sigma > 0 else alpha_star
        alcista   = (len(precios) >= SMA_WINDOW and
                     precios.iloc[-1] > precios.iloc[-SMA_WINDOW:].mean())
        alpha_obj = alpha_vol if alcista else alpha_vol * 0.5
        evaluar_banda(alpha_obj)
        if razon:
            razon += f" | regimen={'alcista' if alcista else 'BAJISTA'} alpha_obj={alpha_obj:.3f}"
        else:
            razon = f"Regimen {'alcista' if alcista else 'BAJISTA'} | dentro de banda DN"

    elif ESTRATEGIA == "E5":
        alpha_vol = min(SIGMA_TARGET / sigma, MAX_LEVERAGE) if sigma > 0 else alpha_star
        alcista   = (len(precios) >= SMA_WINDOW and
                     precios.iloc[-1] > precios.iloc[-SMA_WINDOW:].mean())
        dd        = (valor_cartera / peak_valor) - 1 if peak_valor > 0 else 0
        defensivo = estado.get("en_modo_defensivo", False)
        if   not defensivo and dd < DD_ENTRY_THRESHOLD:
            defensivo = True
        elif defensivo     and dd > DD_EXIT_THRESHOLD:
            defensivo = False
        estado["en_modo_defensivo"] = defensivo
        alpha_base = alpha_vol if alcista else alpha_vol * 0.5
        alpha_obj  = alpha_base * 0.5 if defensivo else alpha_base
        evaluar_banda(alpha_obj)
        if not razon:
            razon = f"DD={dd:.1%} defensivo={defensivo} alpha_obj={alpha_obj:.3f}"

    return {
        "decision":       decision,
        "alpha_nuevo":    round(alpha_nuevo, 4),
        "alpha_anterior": round(alpha_actual, 4),
        "razon":          razon,
        "banda_inf":      round(banda_inf, 4),
        "banda_sup":      round(banda_sup, 4),
        "alpha_star":     round(alpha_star, 4),
    }


# ===========================================================================
# ESTADO
# ===========================================================================

def cargar_estado():
    if ESTADO_FILE.exists():
        with open(ESTADO_FILE) as f:
            return json.load(f)
    return {
        "iniciado": False, "fecha_inicio": None,
        "capital_inicial": CAPITAL_INICIAL,
        "valor_cartera": CAPITAL_INICIAL,
        "cash": CAPITAL_INICIAL,
        "participaciones": 0.0,
        "alpha_actual": 0.0,
        "peak_valor": CAPITAL_INICIAL,
        "en_modo_defensivo": False,
        "n_operaciones": 0,
        "costes_acumulados": 0.0,
    }

def guardar_estado(estado):
    with open(ESTADO_FILE, "w") as f:
        json.dump(estado, f, indent=2, default=str)


# ===========================================================================
# EXCEL OPERATIVA
# ===========================================================================

def generar_excel(fecha, decision, precio, alpha_ant, alpha_nuevo,
                  valor_cartera, importe, coste, n_ops, costes_acum):
    """
    Genera o actualiza el archivo Operativa_GrupoX.xlsx con el formato AFI.

    Formato exacto del template:
        ID | Cantidad | Precio | CT | Precio Ejecutado

    Reglas por decision:
        COMPRAR : Cantidad = participaciones compradas (>0)
                  Precio   = Last Price
                  CT       = lambda del ETF
                  Precio Ejecutado = Precio * (1 + CT)
        VENDER  : Cantidad = participaciones vendidas (<0, negativo)
                  Precio   = Last Price
                  CT       = lambda del ETF
                  Precio Ejecutado = Precio * (1 - CT)
        MANTENER: Cantidad=0, Precio=0, CT=CT, Precio Ejecutado=0

    El archivo es ACUMULATIVO: se añade una fila nueva cada dia.
    Si no existe se crea con las cabeceras del template.
    """
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    # Un unico archivo acumulativo por grupo
    path = OUTPUT_DIR / f"Operativa_Grupo{GRUPO}.xlsx"

    ct    = CT.get(ETF_TICKER, 0.001)
    bold  = Font(name="Calibri", bold=True, size=11)
    norm  = Font(name="Calibri", size=11)
    brd   = Border(left=Side(style="thin"),  right=Side(style="thin"),
                   top=Side(style="thin"),   bottom=Side(style="thin"))
    ctr   = Alignment(horizontal="center", vertical="center")

    # Calcular valores de la fila segun la decision
    if decision == "COMPRAR":
        # Participaciones compradas = importe / precio ejecutado
        precio_ejec = round(precio * (1 + ct), 6)
        cantidad    = round(importe / precio_ejec, 4) if precio_ejec > 0 else 0.0
        precio_fila = round(precio, 4)
    elif decision == "VENDER":
        # Participaciones vendidas = negativo
        precio_ejec = round(precio * (1 - ct), 6)
        cantidad    = -round(importe / precio_ejec, 4) if precio_ejec > 0 else 0.0
        precio_fila = round(precio, 4)
    else:  # MANTENER
        cantidad    = 0
        precio_fila = 0
        precio_ejec = 0

    nueva_fila = [ETF_TICKER, cantidad, precio_fila, ct, precio_ejec]

    # Si el archivo no existe, crearlo con las cabeceras del template
    if not path.exists():
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Operativa"

        # Cabeceras exactas del template AFI
        cabeceras = ["ID", "Cantidad", "Precio", "CT", "Precio Ejecutado"]
        for col, txt in enumerate(cabeceras, 1):
            cell = ws.cell(row=1, column=col, value=txt)
            cell.font      = bold
            cell.border    = brd
            cell.alignment = ctr

        # Anchos de columna
        for col, w in zip("ABCDE", [14, 14, 14, 12, 18]):
            ws.column_dimensions[col].width = w
    else:
        wb = openpyxl.load_workbook(path)
        ws = wb.active

    # Añadir la fila nueva al final
    next_row = ws.max_row + 1
    formatos = [None, "#,##0.0000", "#,##0.0000", "0.000000%", "#,##0.0000"]

    for col, (val, fmt) in enumerate(zip(nueva_fila, formatos), 1):
        cell = ws.cell(row=next_row, column=col, value=val)
        cell.font      = norm
        cell.border    = brd
        cell.alignment = ctr
        if fmt:
            cell.number_format = fmt

    wb.save(path)
    return path


# ===========================================================================
# ENVIO EMAIL
# ===========================================================================

def enviar_email(fecha, decision, resultado, valor_cartera,
                 importe, coste, params, excel_path, n_ops, costes_acum):
    asunto = (f"[GQ_2026] - Grupo{GRUPO} | "
              f"{fecha.strftime('%d/%m/%Y')} | {decision}")
    retorno = (valor_cartera / CAPITAL_INICIAL) - 1

    detalle_op = ""
    if decision != "MANTENER":
        detalle_op = f"""
Detalle de la operacion:
  Alpha anterior:     {resultado['alpha_anterior']:.4f}
  Alpha nuevo:        {resultado['alpha_nuevo']:.4f}
  Importe:            {importe:,.2f} EUR
  Coste transaccion:  {coste:,.2f} EUR
  Razon:              {resultado['razon']}
"""

    cuerpo = f"""Estimado equipo AFI,

Le remitimos la operativa del {fecha.strftime('%d/%m/%Y')} — Grupo {GRUPO}.

ETF:        {ETF_TICKER}
Estrategia: {ESTRATEGIA}
Decision:   {decision}
{detalle_op}
Estado de la cartera:
  Valor cartera:      {valor_cartera:,.2f} EUR
  Retorno acumulado:  {retorno:+.2%}
  Alpha actual:       {resultado['alpha_nuevo']:.4f}
  Banda DN:           [{resultado['banda_inf']:.4f}, {resultado['banda_sup']:.4f}]
  N operaciones:      {n_ops}
  Costes acumulados:  {costes_acum:,.2f} EUR

Parametros del modelo estimados hoy:
  mu:    {params['mu']:.2%}
  sigma: {params['sigma']:.2%}
  alpha* Merton:      {params['alpha_star']:.4f}
  Tasa libre riesgo:  {params['rf']:.2%}

{"Se adjunta Operativa_Grupo" + str(GRUPO) + ".xlsx" if decision != "MANTENER"
  else "No se adjunta operativa: no hay operacion hoy."}

Atentamente,
Grupo {GRUPO} — Master en Finanzas Cuantitativas AFI
"""

    msg           = MIMEMultipart()
    msg["From"]   = EMAIL_REMITENTE
    msg["To"]     = EMAIL_DESTINO
    msg["Subject"] = asunto
    if EMAIL_CC:
        msg["Cc"] = EMAIL_CC
    msg.attach(MIMEText(cuerpo, "plain", "utf-8"))

    if decision != "MANTENER" and excel_path and excel_path.exists():
        with open(excel_path, "rb") as f:
            part = MIMEBase("application", "octet-stream")
            part.set_payload(f.read())
        encoders.encode_base64(part)
        part.add_header("Content-Disposition",
                        f"attachment; filename={excel_path.name}")
        msg.attach(part)

    destinatarios = [EMAIL_DESTINO] + ([EMAIL_CC] if EMAIL_CC else [])
    try:
        with smtplib.SMTP("smtp-mail.outlook.com", 587) as server:
            server.starttls()
            server.login(EMAIL_REMITENTE, EMAIL_PASSWORD)
            server.sendmail(EMAIL_REMITENTE, destinatarios, msg.as_string())
        print(f"  Email enviado a {EMAIL_DESTINO}")
        return True
    except Exception as e:
        print(f"  ERROR al enviar email: {e}")
        print(f"  Envia manualmente: {excel_path}")
        return False


# ===========================================================================
# LOG CSV
# ===========================================================================

def registrar_log(fecha, precio, params, resultado, valor_cartera, coste):
    fila = {
        "fecha":          fecha,
        "precio_etf":     round(precio, 4),
        "mu":             round(params["mu"], 4),
        "sigma":          round(params["sigma"], 4),
        "alpha_star":     round(params["alpha_star"], 4),
        "banda_inf":      round(params["banda_inf"], 4),
        "banda_sup":      round(params["banda_sup"], 4),
        "alpha_anterior": round(resultado["alpha_anterior"], 4),
        "alpha_nuevo":    round(resultado["alpha_nuevo"], 4),
        "decision":       resultado["decision"],
        "razon":          resultado["razon"],
        "valor_cartera":  round(valor_cartera, 2),
        "coste_eur":      round(coste, 2),
        "retorno_acum":   round((valor_cartera / CAPITAL_INICIAL) - 1, 6),
    }
    pd.DataFrame([fila]).to_csv(
        LOG_FILE, mode="a",
        header=not LOG_FILE.exists(),
        index=False
    )


# ===========================================================================
# EXCEL ACUMULATIVO
# ===========================================================================

def actualizar_historial(fecha, precio, params, resultado, valor_cartera,
                          importe, coste, n_ops, costes_acum):
    """Añade una fila al Excel acumulativo Historial_GrupoX.xlsx"""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    historial_path = OUTPUT_DIR / f"Historial_Grupo{GRUPO}.xlsx"

    # Fila nueva
    retorno = round((valor_cartera / CAPITAL_INICIAL) - 1, 6)
    fila = {
        "Fecha":            fecha,
        "Precio ETF":       round(precio, 4),
        "Decision":         resultado["decision"],
        "Alpha anterior":   round(resultado["alpha_anterior"], 4),
        "Alpha nuevo":      round(resultado["alpha_nuevo"], 4),
        "Banda inf":        round(resultado["banda_inf"], 4),
        "Banda sup":        round(resultado["banda_sup"], 4),
        "Alpha* Merton":    round(resultado["alpha_star"], 4),
        "mu":               round(params["mu"], 4),
        "sigma":            round(params["sigma"], 4),
        "Importe (EUR)":    round(importe, 2),
        "Coste (EUR)":      round(coste, 2),
        "Valor cartera":    round(valor_cartera, 2),
        "Retorno acum.":    retorno,
        "N operaciones":    n_ops,
        "Costes acum.":     round(costes_acum, 2),
        "Razon":            resultado["razon"],
    }

    # Cargar historial existente o crear nuevo
    if historial_path.exists():
        wb = openpyxl.load_workbook(historial_path)
        ws = wb.active
        next_row = ws.max_row + 1
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Historial"
        next_row = 2

        # Estilos cabecera
        hdr_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        hdr_fill = PatternFill("solid", start_color="1F4E79")
        hdr_aln  = Alignment(horizontal="center", vertical="center")
        brd      = Border(left=Side(style="thin"),  right=Side(style="thin"),
                          top=Side(style="thin"),   bottom=Side(style="thin"))

        # Título
        ws.merge_cells(f"A1:{chr(64+len(fila))}1")
        title_cell = ws.cell(row=1, column=1,
            value=f"HISTORIAL SIMULACION — GRUPO {GRUPO} | ETF: {ETF_TICKER} | Estrategia: {ESTRATEGIA}")
        title_cell.font = hdr_font
        title_cell.fill = hdr_fill
        title_cell.alignment = hdr_aln
        ws.row_dimensions[1].height = 22

        # Cabeceras columnas
        for col, nombre in enumerate(fila.keys(), 1):
            c = ws.cell(row=2, column=col, value=nombre)
            c.font  = hdr_font
            c.fill  = hdr_fill
            c.alignment = hdr_aln
            c.border = brd
        ws.row_dimensions[2].height = 18
        next_row = 3

    # Estilos fila de datos
    norm_font = Font(name="Arial", size=10)
    brd       = Border(left=Side(style="thin"),  right=Side(style="thin"),
                       top=Side(style="thin"),   bottom=Side(style="thin"))
    gfil = PatternFill("solid", start_color="E2EFDA")
    rfil = PatternFill("solid", start_color="FCE4D6")
    bfil = PatternFill("solid", start_color="DDEBF7")

    decision = resultado["decision"]
    row_fill = gfil if decision == "COMPRAR" else rfil if decision == "VENDER" else None

    formatos = {
        "Fecha":          "DD/MM/YYYY",
        "Precio ETF":     "#,##0.0000",
        "Alpha anterior": "0.0000",
        "Alpha nuevo":    "0.0000",
        "Banda inf":      "0.0000",
        "Banda sup":      "0.0000",
        "Alpha* Merton":  "0.0000",
        "mu":             "0.00%",
        "sigma":          "0.00%",
        "Importe (EUR)":  "#,##0.00",
        "Coste (EUR)":    "#,##0.00",
        "Valor cartera":  "#,##0.00",
        "Retorno acum.":  "0.00%",
        "Costes acum.":   "#,##0.00",
    }

    for col, (nombre, valor) in enumerate(fila.items(), 1):
        c = ws.cell(row=next_row, column=col, value=valor)
        c.font      = norm_font
        c.border    = brd
        c.alignment = Alignment(horizontal="center", vertical="center")
        if row_fill and nombre == "Decision":
            c.fill = row_fill
        if nombre in formatos:
            c.number_format = formatos[nombre]
        if nombre == "Retorno acum.":
            c.fill = gfil if retorno >= 0 else rfil

    # Anchos de columna (solo si es archivo nuevo)
    if next_row == 3:
        anchos = [12, 12, 11, 14, 11, 10, 10, 13, 8, 8, 14, 12, 14, 13, 13, 12, 35]
        for col, ancho in enumerate(anchos, 1):
            ws.column_dimensions[chr(64 + col)].width = ancho

    ws.row_dimensions[next_row].height = 16
    wb.save(historial_path)
    return historial_path


# ===========================================================================
# MAIN
# ===========================================================================

def es_dia_habil():
    return date.today().weekday() < 5


def main():
    global estado

    print("\n" + "="*65)
    print(f"  SIMULACION DIARIA — {ESTRATEGIA} sobre {ETF_TICKER}")
    print(f"  {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}")
    print("="*65)

    if not es_dia_habil():
        print("\n  Hoy es fin de semana. No se ejecuta.")
        return

    estado = cargar_estado()

    # Descargar datos
    print(f"\n[1/4] Descargando {ETF_TICKER}...")
    datos = yf.download(ETF_TICKER, start="2008-01-01", progress=False)
    if datos.empty:
        print("  ERROR: no se pudieron descargar datos.")
        sys.exit(1)
    precios    = datos["Close"].squeeze().dropna()
    precio_hoy = float(precios.iloc[-1])
    fecha_hoy  = precios.index[-1].date()
    print(f"  {fecha_hoy}: {precio_hoy:.4f} EUR")

    try:
        rf_data  = yf.download("ECBDFR", start="2020-01-01", progress=False)
        rf_anual = float(rf_data["Close"].iloc[-1]) / 100
    except Exception:
        rf_anual = 0.02
    print(f"  Tasa BCE: {rf_anual:.2%}")

    # Parametros
    print(f"\n[2/4] Estimando parametros...")
    params = estimar_parametros(precios, rf_anual)
    if params is None:
        print("  ERROR: datos insuficientes.")
        sys.exit(1)
    params["precio_hoy"] = precio_hoy
    print(f"  mu={params['mu']:.2%} sigma={params['sigma']:.2%} "
          f"a*={params['alpha_star']:.4f} "
          f"banda=[{params['banda_inf']:.4f},{params['banda_sup']:.4f}]")

    # Valor cartera
    participaciones = estado["participaciones"]
    cash            = estado["cash"]
    valor_cartera   = cash + participaciones * precio_hoy
    peak_valor      = max(estado["peak_valor"], valor_cartera)
    estado["peak_valor"] = peak_valor

    print(f"\n[3/4] Cartera: {valor_cartera:,.2f} EUR | "
          f"alpha={estado['alpha_actual']:.4f} | "
          f"DD={(valor_cartera/peak_valor - 1):.2%}")

    # Decision
    print(f"\n[4/4] Estrategia {ESTRATEGIA}...")
    resultado   = decidir_alpha(params, precios, estado["alpha_actual"],
                                peak_valor, valor_cartera, estado)
    decision    = resultado["decision"]
    alpha_nuevo = resultado["alpha_nuevo"]
    importe     = 0.0
    coste       = 0.0
    ct          = CT.get(ETF_TICKER, 0.001)

    if decision != "MANTENER":
        diferencia = (alpha_nuevo * valor_cartera) - (participaciones * precio_hoy)
        importe    = abs(diferencia)
        coste      = importe * ct
        if decision == "COMPRAR":
            p_ejec = precio_hoy * (1 + ct)
            nuevas = diferencia / p_ejec
            participaciones += nuevas
            cash            -= nuevas * p_ejec
        else:
            p_ejec = precio_hoy * (1 - ct)
            vender = importe / p_ejec
            participaciones = max(0.0, participaciones - vender)
            cash            += vender * p_ejec

        valor_cartera               = cash + participaciones * precio_hoy
        estado["n_operaciones"]    += 1
        estado["costes_acumulados"] += coste

    estado.update({
        "valor_cartera":   valor_cartera,
        "cash":            cash,
        "participaciones": participaciones,
        "alpha_actual":    alpha_nuevo,
        "iniciado":        True,
        "fecha_inicio":    estado["fecha_inicio"] or str(fecha_hoy),
    })

    n_ops       = estado["n_operaciones"]
    costes_acum = estado["costes_acumulados"]

    emoji = ("COMPRAR" if decision == "COMPRAR" else
             "VENDER"  if decision == "VENDER"  else "MANTENER")
    print(f"\n  DECISION: {emoji}")
    print(f"  Alpha: {resultado['alpha_anterior']:.4f} -> {alpha_nuevo:.4f}")
    print(f"  Razon: {resultado['razon']}")
    print(f"  Valor cartera:  {valor_cartera:,.2f} EUR")
    print(f"  Retorno acum.:  {((valor_cartera/CAPITAL_INICIAL)-1):+.2%}")
    if decision != "MANTENER":
        print(f"  Importe:        {importe:,.2f} EUR")
        print(f"  Coste:          {coste:,.2f} EUR")

    guardar_estado(estado)
    registrar_log(fecha_hoy, precio_hoy, params, resultado, valor_cartera, coste)

    # Excel diario (para enviar a AFI)
    excel_path = generar_excel(
        fecha_hoy, decision, precio_hoy,
        resultado["alpha_anterior"], alpha_nuevo,
        valor_cartera, importe, coste, n_ops, costes_acum
    )
    print(f"\n  Excel diario:      {excel_path.name}")

    # Excel acumulativo (historial propio)
    hist_path = actualizar_historial(
        fecha_hoy, precio_hoy, params, resultado, valor_cartera,
        importe, coste, n_ops, costes_acum
    )
    print(f"  Excel historial:   {hist_path.name}")

    print(f"  Enviando email...")
    enviar_email(fecha_hoy, decision, resultado, valor_cartera,
                 importe, coste, params, excel_path, n_ops, costes_acum)

    print("="*65 + "\n")


if __name__ == "__main__":
    main()