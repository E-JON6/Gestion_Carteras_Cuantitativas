"""Build and email daily trade instructions with an Excel attachment."""

from __future__ import annotations

import smtplib
import ssl
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path
from typing import Any, Mapping, Sequence

from openpyxl import load_workbook
from openpyxl.styles import NamedStyle

from services.settings import TradeEmailSettings

SMTP_HOST = "smtp.gmail.com"
SMTP_PORT = 465
WORKSHEET_NAME = "Operativa"
REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_TEMPLATE_PATH = REPO_ROOT / "docs" / "Operativa_GrupoX.xlsx"
DEFAULT_EXAMPLE_PATH = REPO_ROOT / "docs" / "Ejemplo_CT.xlsx"
DEFAULT_OUTPUT_DIR = REPO_ROOT / "outputs"
EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass(frozen=True)
class TradeInstruction:
    id: str
    quantity: float
    price: float
    ct: float


def _coerce_float(value: Any, field_name: str, row_number: int) -> float:
    if isinstance(value, bool):
        raise ValueError(f"Trade #{row_number}: '{field_name}' must be numeric, got boolean")

    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(
            f"Trade #{row_number}: '{field_name}' must be numeric, got {value!r}"
        ) from exc


def validate_trades(
    trades: Sequence[TradeInstruction | Mapping[str, Any]],
) -> list[TradeInstruction]:
    """Validate and normalize raw trade dictionaries."""
    if not trades:
        raise ValueError("At least one trade is required")

    validated: list[TradeInstruction] = []

    for row_number, trade in enumerate(trades, start=1):
        if isinstance(trade, TradeInstruction):
            normalized = trade
        elif isinstance(trade, Mapping):
            missing = [key for key in ("id", "quantity", "price", "ct") if key not in trade]
            if missing:
                raise ValueError(
                    f"Trade #{row_number}: missing required fields: {', '.join(missing)}"
                )

            trade_id = str(trade["id"]).strip() if trade["id"] is not None else ""
            quantity = _coerce_float(trade["quantity"], "quantity", row_number)
            price = _coerce_float(trade["price"], "price", row_number)
            ct = _coerce_float(trade["ct"], "ct", row_number)
            normalized = TradeInstruction(
                id=trade_id,
                quantity=quantity,
                price=price,
                ct=ct,
            )
        else:
            raise TypeError(
                f"Trade #{row_number}: expected a mapping or TradeInstruction, got {type(trade).__name__}"
            )

        if not normalized.id:
            raise ValueError(f"Trade #{row_number}: 'id' must be non-empty")
        if normalized.quantity == 0:
            raise ValueError(f"Trade #{row_number}: 'quantity' must be non-zero")
        if normalized.price <= 0:
            raise ValueError(f"Trade #{row_number}: 'price' must be positive")
        if normalized.ct < 0:
            raise ValueError(f"Trade #{row_number}: 'ct' must be non-negative")

        validated.append(normalized)

    return validated


def _load_template_workbook(template_path: Path):
    if not template_path.exists():
        raise FileNotFoundError(f"Excel template not found: {template_path}")

    workbook = load_workbook(template_path)
    if WORKSHEET_NAME not in workbook.sheetnames:
        raise ValueError(
            f"Worksheet '{WORKSHEET_NAME}' was not found in template: {template_path}"
        )

    return workbook


def _extract_ct_style(example_path: Path) -> tuple[NamedStyle | None, str]:
    """Read the CT formatting from the example workbook when available."""
    if not example_path.exists():
        return None, "0.00%"

    workbook = load_workbook(example_path)
    sheet = workbook.worksheets[0]
    ct_cell = sheet["D2"]
    return copy(ct_cell._style), ct_cell.number_format or "0.00%"


def _clear_existing_trade_rows(worksheet) -> None:
    if worksheet.max_row > 1:
        worksheet.delete_rows(2, worksheet.max_row - 1)


def _formula_for_row(quantity: float, row_number: int) -> str:
    if quantity > 0:
        return f"=C{row_number}*(1+D{row_number})"
    if quantity < 0:
        return f"=C{row_number}*(1-D{row_number})"
    raise ValueError(f"Trade row {row_number}: quantity cannot be zero")


def build_operativa_attachment(
    trades: Sequence[TradeInstruction | Mapping[str, Any]],
    template_path: str | Path = DEFAULT_TEMPLATE_PATH,
    output_dir: str | Path = DEFAULT_OUTPUT_DIR,
    *,
    example_path: str | Path = DEFAULT_EXAMPLE_PATH,
) -> Path:
    """Create a filled Operativa workbook and return the saved file path."""
    validated_trades = validate_trades(trades)
    template = Path(template_path)
    output_directory = Path(output_dir)
    example = Path(example_path)

    workbook = _load_template_workbook(template)
    worksheet = workbook[WORKSHEET_NAME]
    _clear_existing_trade_rows(worksheet)

    ct_style, ct_number_format = _extract_ct_style(example)

    for index, trade in enumerate(validated_trades, start=2):
        worksheet.cell(row=index, column=1, value=trade.id)
        worksheet.cell(row=index, column=2, value=trade.quantity)
        worksheet.cell(row=index, column=3, value=trade.price)

        ct_cell = worksheet.cell(row=index, column=4, value=trade.ct)
        if ct_style is not None:
            ct_cell._style = copy(ct_style)
        ct_cell.number_format = ct_number_format

        formula = _formula_for_row(trade.quantity, index)
        worksheet.cell(row=index, column=5, value=formula)

    output_directory.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = output_directory / f"{template.stem}_{timestamp}.xlsx"
    workbook.save(output_path)
    return output_path


def send_email_with_attachment(
    subject: str,
    body: str,
    attachment_path: str | Path,
    to: Sequence[str],
    *,
    cc: Sequence[str] | None = None,
    sender_email: str,
    app_password: str,
) -> None:
    """Send an email with the generated Excel workbook attached."""
    attachment = Path(attachment_path)
    if not attachment.exists():
        raise FileNotFoundError(f"Attachment not found: {attachment}")

    recipients_to = [recipient.strip() for recipient in to if recipient.strip()]
    recipients_cc = [recipient.strip() for recipient in (cc or []) if recipient.strip()]
    if not recipients_to:
        raise ValueError("At least one 'to' recipient is required")

    message = EmailMessage()
    message["From"] = sender_email
    message["To"] = ", ".join(recipients_to)
    if recipients_cc:
        message["Cc"] = ", ".join(recipients_cc)
    message["Subject"] = subject
    message.set_content(body)

    attachment_bytes = attachment.read_bytes()
    maintype, subtype = EXCEL_MIME.split("/", maxsplit=1)
    message.add_attachment(
        attachment_bytes,
        maintype=maintype,
        subtype=subtype,
        filename=attachment.name,
    )

    all_recipients = recipients_to + recipients_cc
    context = ssl.create_default_context()

    try:
        with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, context=context) as smtp:
            smtp.login(sender_email, app_password)
            smtp.send_message(message, to_addrs=all_recipients)
    except smtplib.SMTPException as exc:
        raise RuntimeError(f"Failed to send Gmail SMTP message: {exc}") from exc
    except OSError as exc:
        raise RuntimeError(f"SMTP connection failed: {exc}") from exc


def send_daily_trade_email(
    trades: Sequence[TradeInstruction | Mapping[str, Any]],
    template_path: str | Path = DEFAULT_TEMPLATE_PATH,
    output_dir: str | Path = DEFAULT_OUTPUT_DIR,
    subject: str | None = None,
    body: str | None = None,
) -> dict[str, Any]:
    """Validate trades, generate the workbook, and send the email."""
    settings = TradeEmailSettings.from_env()
    validated_trades = validate_trades(trades)
    attachment_path = build_operativa_attachment(
        validated_trades,
        template_path=template_path,
        output_dir=output_dir,
    )

    resolved_subject = subject or settings.default_subject
    resolved_body = body or settings.default_body

    send_email_with_attachment(
        resolved_subject,
        resolved_body,
        attachment_path,
        settings.to,
        cc=settings.cc,
        sender_email=settings.smtp_email,
        app_password=settings.app_password,
    )

    return {
        "status": "sent",
        "attachment_path": str(attachment_path),
        "to": ", ".join(settings.to),
        "cc": ", ".join(settings.cc),
        "subject": resolved_subject,
        "num_trades": len(validated_trades),
    }
